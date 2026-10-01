"""Test layer tool PDF.

    .venv/bin/python -m pytest tests/test_tools.py -q
"""

from __future__ import annotations

import os
import sys

import pymupdf
import pytest

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from pdfocr import tools  # noqa: E402
from pdfocr.tools.base import (  # noqa: E402
    INPUT_ANY,
    INPUT_IMAGE,
    INPUT_NONE,
    INPUT_OFFICE,
    INPUT_PDF,
    INPUT_PDF_MULTI,
    Param,
    ToolError,
    collect_widgets,
    parse_page_range,
    run_spec,
)
from pdfocr.tools.edit import parse_field_values  # noqa: E402


# ------------------------------------------------------------------- fixtures

def make_pdf(path: str, pages: int = 4, text: str = "Rahasia XYZ") -> str:
    doc = pymupdf.open()
    for i in range(pages):
        page = doc.new_page(width=595, height=842)
        page.insert_text((72, 120), f"Halaman {i + 1} {text}", fontsize=22, fontname="hebo")
    doc.save(path)
    doc.close()
    return path


def make_form_pdf(path: str) -> str:
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_text((72, 90), "Formulir", fontsize=18, fontname="hebo")
    for index, name in enumerate(("nama", "nim", "kota")):
        y = 140 + index * 50
        page.insert_text((72, y), f"{name}:", fontsize=11)
        widget = pymupdf.Widget()
        widget.field_name = name
        widget.field_type = pymupdf.PDF_WIDGET_TYPE_TEXT
        widget.rect = pymupdf.Rect(180, y - 14, 420, y + 6)
        widget.field_value = ""
        page.add_widget(widget)
    doc.save(path)
    doc.close()
    return path


@pytest.fixture
def sample(tmp_path) -> str:
    return make_pdf(str(tmp_path / "sample.pdf"))


def run(tool_id: str, inputs, params=None, out_path=None, tmp_path=None):
    spec = tools.get(tool_id)
    merged = {p.key: p.default for p in spec.params}
    merged.update(params or {})
    if out_path is None:
        if spec.multi_output or not spec.extension:
            out_path = str(tmp_path / f"out_{tool_id}")
        else:
            out_path = str(tmp_path / f"out_{tool_id}{spec.extension}")
    return run_spec(spec, inputs, merged, out_path), out_path


# -------------------------------------------------------------------- registry

def test_registry_ids_unique():
    ids = [spec.id for spec in tools.ALL_SPECS]
    assert len(ids) == len(set(ids))


def test_registry_specs_are_wellformed():
    valid = {
        INPUT_NONE, INPUT_PDF, INPUT_PDF_MULTI,
        INPUT_IMAGE, INPUT_OFFICE, INPUT_ANY,
    }
    for spec in tools.ALL_SPECS:
        assert spec.id and spec.label and spec.group, spec.id
        assert callable(spec.run), spec.id
        assert spec.input_kind in valid, spec.id
        assert spec.description, spec.id
        keys = [p.key for p in spec.params]
        assert len(keys) == len(set(keys)), spec.id


def test_registry_groups_match_menu_order():
    for group, specs in tools.grouped().items():
        assert specs, group
        for spec in specs:
            assert spec.group == group


def test_every_group_has_at_least_one_menu_item():
    grouped = tools.grouped()
    native = {"merge", "remove", "ocr", "rotate"}
    for name in tools.MENU_ORDER:
        remaining = [s for s in grouped.get(name, []) if s.id not in native]
        assert remaining, f"{name} punya menu kosong"


def test_get_unknown_tool_raises():
    with pytest.raises(ToolError):
        tools.get("tidak-ada")


# ------------------------------------------------------------------- base api

@pytest.mark.parametrize(
    "text,expected",
    [
        ("", [0, 1, 2]),
        ("all", [0, 1, 2]),
        ("1", [0]),
        ("2-3", [1, 2]),
        ("1,3", [0, 2]),
        ("2-", [1, 2]),
        ("3-1", [0, 1, 2]),
        ("1-99", [0, 1, 2]),
    ],
)
def test_parse_page_range(text, expected):
    assert parse_page_range(text, 3) == expected


@pytest.mark.parametrize("text", ["abc", "1-xyz", "x-y", "1,,,z"])
def test_parse_page_range_rejects_garbage(text):
    with pytest.raises(ToolError):
        parse_page_range(text, 3)


def test_parse_page_range_open_edges_are_all():
    # "0-" dan "-" diartikan "dari halaman pertama".
    assert parse_page_range("0-", 3) == [0, 1, 2]
    assert parse_page_range("-", 3) == [0, 1, 2]


def test_parse_page_range_out_of_range():
    with pytest.raises(ToolError):
        parse_page_range("99", 3)


def test_param_parse_types():
    assert Param("k", "K", "int", "0").parse(" 7 ") == 7
    assert Param("k", "K", "float", "0").parse("1.5") == 1.5
    assert Param("k", "K", "bool", "").parse("1") is True
    assert Param("k", "K", "bool", "1").parse("") is False
    assert Param("k", "K", "text", "").parse(" halo ") == " halo "  # spasi dipertahankan


def test_param_int_error_message():
    with pytest.raises(ToolError):
        Param("tebal", "Tebal", "int", "0").parse("abc")


def test_param_choice_labels_map_to_internal():
    p = Param("mode", "Mode", "choice", "a", choices=("a", "b"), labels=("A", "B"))
    assert p.display_choices() == ["A", "B"]
    assert p.internal_value("B") == "b"
    assert p.internal_value("ngawur") == "a"


def test_parse_field_values():
    assert parse_field_values("a=1;b=2") == {"a": "1", "b": "2"}
    assert parse_field_values("a=1\nb=2") == {"a": "1", "b": "2"}
    assert parse_field_values("") == {}
    assert parse_field_values("salah") == {}
    assert parse_field_values("a=x=y") == {"a": "x=y"}


def test_run_spec_reports_missing_input():
    spec = tools.get("split")
    with pytest.raises(ToolError):
        run_spec(spec, [], {}, "/tmp/nt/anything")


def test_outcome_message_includes_all_parts():
    from pdfocr.tools.base import ToolOutcome

    text = ToolOutcome(
        output="/tmp/a.pdf", summary="Selesai.", warnings=["w1"], items=["/tmp/b.pdf"]
    ).message()
    assert "Selesai." in text
    assert "a.pdf" in text
    assert "w1" in text
    assert "b.pdf" in text


# --------------------------------------------------------------------- merge

def test_merge_preserves_page_count(sample, tmp_path):
    outcome, out = run("merge", [sample], out_path=str(tmp_path / "m.pdf"), tmp_path=tmp_path)
    doc = pymupdf.open(out)
    assert doc.page_count == 4
    doc.close()


def test_remove_pages(sample, tmp_path):
    outcome, out = run("remove", [sample], {"ranges": "1,3"}, tmp_path=tmp_path)
    doc = pymupdf.open(out)
    assert doc.page_count == 2
    assert "Halaman 2" in doc[0].get_text()
    assert "Halaman 4" in doc[1].get_text()
    doc.close()


def test_remove_all_pages_rejected(sample, tmp_path):
    with pytest.raises(ToolError):
        run("remove", [sample], {"ranges": "1-4"}, tmp_path=tmp_path)


def test_extract_pages(sample, tmp_path):
    outcome, out = run("extract", [sample], {"ranges": "2-3"}, tmp_path=tmp_path)
    doc = pymupdf.open(out)
    assert doc.page_count == 2
    doc.close()


def test_split_each_page(sample, tmp_path):
    outcome, out = run("split", [sample], {"mode": "each"}, tmp_path=tmp_path)
    assert len(os.listdir(out)) == 4


def test_split_by_ranges(sample, tmp_path):
    outcome, out = run("split", [sample], {"mode": "ranges", "ranges": "1-2,3-4"},
                       tmp_path=tmp_path)
    assert len(os.listdir(out)) == 2


def test_extract_separate(sample, tmp_path):
    outcome, out = run("extract", [sample], {"ranges": "1,2", "separate": "1"},
                       tmp_path=tmp_path)
    assert len(os.listdir(out)) == 2


# ---------------------------------------------------------------- edit tools

def test_rotate_applies_to_selected_pages(sample, tmp_path):
    outcome, out = run("rotate", [sample], {"angle": "90", "ranges": "1"}, tmp_path=tmp_path)
    doc = pymupdf.open(out)
    assert doc[0].rotation == 90
    assert doc[1].rotation == 0
    doc.close()


def test_rotate_rejects_bad_angle(sample, tmp_path):
    with pytest.raises(ToolError):
        run("rotate", [sample], {"angle": "37"}, tmp_path=tmp_path)


def test_watermark_arbitrary_angle_and_text(sample, tmp_path):
    outcome, out = run(
        "watermark", [sample],
        {"text": "RAHASIA", "rotation": "37", "layout": "center", "opacity": "0.3"},
        tmp_path=tmp_path,
    )
    doc = pymupdf.open(out)
    assert "RAHASIA" in doc[0].get_text()
    doc.close()


def test_watermark_rejects_empty_text(sample, tmp_path):
    with pytest.raises(ToolError):
        run("watermark", [sample], {"text": "  "}, tmp_path=tmp_path)


def test_add_page_numbers(sample, tmp_path):
    run("numbers", [sample], {"ranges": "1-2", "start": "1"}, tmp_path=tmp_path)
    doc = pymupdf.open(str(tmp_path / "out_numbers.pdf"))
    assert "1" in doc[0].get_text()
    doc.close()


def test_crop_reduces_media_box(sample, tmp_path):
    _, out = run("crop", [sample], {"mode": "margin", "margin": "50"}, tmp_path=tmp_path)
    doc = pymupdf.open(out)
    assert doc[0].rect.width == pytest.approx(495, abs=1)
    doc.close()


def test_annotate_requires_text(sample, tmp_path):
    with pytest.raises(ToolError):
        run("annotate", [sample], {"text": ""}, tmp_path=tmp_path)


def test_annotate_writes_note(sample, tmp_path):
    _, out = run("annotate", [sample], {"text": "Catatan uji", "color": "merah", "page": "1"},
                 tmp_path=tmp_path)
    doc = pymupdf.open(out)
    # Catatan adalah anotasi PDF, jadi tidak muncul di teks halaman.
    annots = list(doc[0].annots())
    assert annots, "tidak ada anotasi"
    assert "Catatan uji" in (annots[0].info or {}).get("content", "")
    doc.close()


# -------------------------------------------------------------------- forms

def test_forms_rejects_non_form_pdf(sample, tmp_path):
    with pytest.raises(ToolError):
        run("forms", [sample], {"values": "a=1"}, tmp_path=tmp_path)


def test_forms_fills_fields(tmp_path):
    source = make_form_pdf(str(tmp_path / "form.pdf"))
    outcome, out = run("forms", [source], {"values": "nama=Budi;nim=12345"},
                       tmp_path=tmp_path)
    doc = pymupdf.open(out)
    pages, widgets = collect_widgets(doc)
    values = {w.field_name: w.field_value for w in widgets}
    assert values["nama"] == "Budi"
    assert values["nim"] == "12345"
    assert values["kota"] == ""          # tidak diisi
    assert "Budi" in doc[0].get_text()   # nilai tampil di halaman
    doc.close()


def test_forms_flatten_removes_widgets(tmp_path):
    source = make_form_pdf(str(tmp_path / "form.pdf"))
    _, out = run("forms", [source], {"values": "nama=Siti", "flatten": "1"}, tmp_path=tmp_path)
    doc = pymupdf.open(out)
    assert not doc.is_form_pdf
    assert "Siti" in doc[0].get_text()   # nilai tetap ada, jadi konten halaman
    doc.close()


def test_forms_rejects_unknown_field(tmp_path):
    source = make_form_pdf(str(tmp_path / "form.pdf"))
    with pytest.raises(ToolError) as info:
        run("forms", [source], {"values": "entah=1"}, tmp_path=tmp_path)
    assert "entah" in str(info.value)


def test_collect_widgets_keeps_pages_alive(tmp_path):
    """Widget harus masih bisa di-update setelah daftar dibuat."""
    source = make_form_pdf(str(tmp_path / "form.pdf"))
    doc = pymupdf.open(source)
    pages, widgets = collect_widgets(doc)
    assert pages and widgets
    for widget in widgets:
        widget.field_value = "OK"
        widget.update()
    doc.save(str(tmp_path / "filled.pdf"))
    doc.close()

    check = pymupdf.open(str(tmp_path / "filled.pdf"))
    assert all(w.field_value == "OK" for w in collect_widgets(check)[1])
    check.close()


# ----------------------------------------------------------------- security

def test_protect_then_unlock_roundtrip(sample, tmp_path):
    _, locked = run("protect", [sample],
                    {"user_password": "pw123", "owner_password": "ow123"},
                    out_path=str(tmp_path / "locked.pdf"), tmp_path=tmp_path)
    doc = pymupdf.open(locked)
    assert doc.needs_pass
    doc.close()

    # Password salah harus ditolak.
    with pytest.raises(ToolError):
        run("unlock", [locked], {"password": "salah"},
            out_path=str(tmp_path / "x.pdf"), tmp_path=tmp_path)

    outcome, plain = run("unlock", [locked], {"password": "pw123"},
                         out_path=str(tmp_path / "plain.pdf"), tmp_path=tmp_path)
    doc = pymupdf.open(plain)
    assert not doc.needs_pass
    assert doc.page_count == 4
    doc.close()


def test_protect_requires_password(sample, tmp_path):
    with pytest.raises(ToolError):
        run("protect", [sample], {"user_password": "", "owner_password": ""},
            tmp_path=tmp_path)


def test_unlock_unencrypted_is_noop(sample, tmp_path):
    outcome, out = run("unlock", [sample], {"password": ""}, tmp_path=tmp_path)
    assert pymupdf.open(out).page_count == 4


def test_redact_removes_keyword(sample, tmp_path):
    _, out = run("redact", [sample], {"mode": "terms", "terms": "Rahasia"},
                 tmp_path=tmp_path)
    doc = pymupdf.open(out)
    assert not doc[0].search_for("Rahasia")
    assert "Halaman" in doc[0].get_text()   # teks lain tetap ada
    doc.close()


def test_redact_reports_when_nothing_found(sample, tmp_path):
    with pytest.raises(ToolError):
        run("redact", [sample], {"mode": "terms", "terms": "TidakAdaIni"},
            tmp_path=tmp_path)


def test_redact_requires_coordinates(tmp_path):
    source = make_pdf(str(tmp_path / "r.pdf"))
    with pytest.raises(ToolError):
        run("redact", [source], {"mode": "rects", "rects": "1,2,3"}, tmp_path=tmp_path)


def test_compare_detects_difference(tmp_path):
    a = make_pdf(str(tmp_path / "a.pdf"), text="Versi satu")
    b = make_pdf(str(tmp_path / "b.pdf"), text="Versi dua")
    outcome, out = run("compare", [a, b], out_path=str(tmp_path / "diff.pdf"),
                       tmp_path=tmp_path)
    doc = pymupdf.open(out)
    text = "".join(doc[i].get_text() for i in range(doc.page_count))
    assert "Versi dua" in text
    assert doc.page_count >= 1
    doc.close()


def test_compare_needs_two_files(sample, tmp_path):
    with pytest.raises(ToolError):
        run("compare", [sample], out_path=str(tmp_path / "d.pdf"), tmp_path=tmp_path)


def test_sign_image_mode(tmp_path):
    source = make_pdf(str(tmp_path / "s.pdf"), pages=2)
    # Tanda tangan harus berupa gambar raster, bukan PDF.
    stamp_doc = pymupdf.open()
    page = stamp_doc.new_page(width=120, height=60)
    page.insert_text((10, 40), "Tanda", fontsize=20)
    stamp_path = str(tmp_path / "stamp.png")
    stamp_doc[0].get_pixmap(dpi=100).save(stamp_path)
    stamp_doc.close()

    _, out = run("sign", [source], {"mode": "image", "image": stamp_path,
                                    "ranges": "1", "width": "120"},
                 out_path=str(tmp_path / "signed.pdf"), tmp_path=tmp_path)
    doc = pymupdf.open(out)
    assert doc.page_count == 2
    assert doc[0].get_images(), "gambar tanda tangan tidak ada"
    doc.close()


def test_sign_rejects_pdf_as_signature_image(tmp_path):
    source = make_pdf(str(tmp_path / "s.pdf"), pages=1)
    fake = make_pdf(str(tmp_path / "notimage.pdf"), pages=1)
    with pytest.raises(ToolError):
        run("sign", [source], {"mode": "image", "image": fake},
            out_path=str(tmp_path / "s.pdf"), tmp_path=tmp_path)


def test_sign_image_requires_file(sample, tmp_path):
    with pytest.raises(ToolError):
        run("sign", [sample], {"mode": "image", "image": ""},
            out_path=str(tmp_path / "s.pdf"), tmp_path=tmp_path)


# ----------------------------------------------------------------- optimize

def test_repair_handles_damaged_xref(tmp_path):
    raw = open(make_pdf(str(tmp_path / "src.pdf")), "rb").read()
    broken = str(tmp_path / "broken.pdf")
    with open(broken, "wb") as handle:
        handle.write(raw[:-300] + b"\x00" * 300)

    outcome, out = run("repair", [broken], out_path=str(tmp_path / "fixed.pdf"),
                       tmp_path=tmp_path)
    assert pymupdf.open(out).page_count == 4


def test_repair_falls_back_to_ghostscript(tmp_path):
    import shutil

    if not shutil.which("gs"):
        pytest.skip("Ghostscript tidak terpasang")

    dead = str(tmp_path / "dead.pdf")
    with open(dead, "wb") as handle:
        handle.write(b"%PDF-1.7\n" + b"\xde\xad\xbe\xef" * 500)

    with pytest.raises(Exception):
        import pymupdf as _p
        _p.open(dead)

    outcome, out = run("repair", [dead], out_path=str(tmp_path / "rescued.pdf"),
                       tmp_path=tmp_path)
    assert "Ghostscript" in outcome.summary
    assert os.path.exists(out)


def test_compress_runs(sample, tmp_path):
    outcome, out = run("compress", [sample], {"level": "medium"}, tmp_path=tmp_path)
    assert os.path.getsize(out) > 0


def test_optimize_ocr_produces_searchable_pdf(tmp_path):
    source = make_pdf(str(tmp_path / "o.pdf"), pages=1, text="Searchable Document")
    outcome, out = run("ocr", [source], {"dpi": "300", "lang": "eng"},
                       out_path=str(tmp_path / "ocr.pdf"), tmp_path=tmp_path)
    doc = pymupdf.open(out)
    assert doc.page_count == 1
    doc.close()


# ------------------------------------------------------------ convert from pdf

def test_pdf_to_images(tmp_path):
    source = make_pdf(str(tmp_path / "i.pdf"), pages=3)
    outcome, out = run("pdf2jpg", [source], {"dpi": "72", "format": "jpg"}, tmp_path=tmp_path)
    files = [f for f in os.listdir(out) if f.endswith(".jpg")]
    assert len(files) == 3


def test_pdf_to_markdown(tmp_path):
    source = make_pdf(str(tmp_path / "m.pdf"), text="Judul Besar")
    outcome, out = run("pdf2md", [source], out_path=str(tmp_path / "out.md"),
                       tmp_path=tmp_path)
    content = open(out, encoding="utf-8").read()
    assert "Judul Besar" in content
    assert "---" in content


def test_pdf_to_excel_falls_back_to_text(tmp_path):
    source = make_pdf(str(tmp_path / "x.pdf"))
    outcome, out = run("pdf2xls", [source], out_path=str(tmp_path / "out.xlsx"),
                       tmp_path=tmp_path)
    assert os.path.exists(out)


def test_pdf_to_excel_rejects_bad_tolerance(tmp_path):
    source = make_pdf(str(tmp_path / "x.pdf"))
    with pytest.raises(ToolError):
        run("pdf2xls", [source], {"text_tolerance": "abc"},
            out_path=str(tmp_path / "o.xlsx"), tmp_path=tmp_path)


def test_pdf_to_ppt_makes_image_slides(tmp_path):
    source = make_pdf(str(tmp_path / "p.pdf"), pages=3)
    _, out = run("pdf2ppt", [source], out_path=str(tmp_path / "out.pptx"), tmp_path=tmp_path)
    from pptx import Presentation

    assert len(Presentation(out).slides) == 3


def test_pdf_to_pdfa(tmp_path):
    import shutil

    if not shutil.which("gs"):
        pytest.skip("Ghostscript tidak terpasang")
    source = make_pdf(str(tmp_path / "a.pdf"), pages=2)
    outcome, out = run("pdf2pdfa", [source], {"flavour": "2"},
                       out_path=str(tmp_path / "out.pdf"), tmp_path=tmp_path)
    assert "PDF/A" in outcome.summary
    doc = pymupdf.open(out)
    assert doc.page_count == 2
    doc.close()


# -------------------------------------------------------------- convert to pdf

def test_images_to_pdf(tmp_path):
    src = make_pdf(str(tmp_path / "s.pdf"), pages=1)
    png = str(tmp_path / "a.png")
    pymupdf.open(src)[0].get_pixmap(dpi=80).save(png)
    outcome, out = run("jpg2pdf", [png, png], {"dpi": "150"},
                       out_path=str(tmp_path / "imgs.pdf"), tmp_path=tmp_path)
    assert pymupdf.open(out).page_count == 2


def test_html_to_pdf(tmp_path):
    html = tmp_path / "p.html"
    html.write_text("<h1>Judul</h1><p>Isi.</p>", encoding="utf-8")
    outcome, out = run("html2pdf", [str(html)], out_path=str(tmp_path / "h.pdf"),
                       tmp_path=tmp_path)
    assert pymupdf.open(out).page_count == 1


def test_office_to_pdf_word(tmp_path):
    import shutil

    if not shutil.which("soffice"):
        pytest.skip("LibreOffice tidak terpasang")
    from docx import Document

    docx_path = str(tmp_path / "d.docx")
    doc = Document()
    doc.add_paragraph("Halo dari Word")
    doc.save(docx_path)

    outcome, out_dir = run("word2pdf", [docx_path], out_path=str(tmp_path / "conv"),
                           tmp_path=tmp_path)
    produced = [f for f in os.listdir(out_dir) if f.endswith(".pdf")]
    assert produced, os.listdir(out_dir)


def test_office_rejects_wrong_extension(sample, tmp_path):
    with pytest.raises(ToolError):
        run("word2pdf", [sample], out_path=str(tmp_path / "w"), tmp_path=tmp_path)
